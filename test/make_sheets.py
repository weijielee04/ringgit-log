"""Test spreadsheets for the import flow: the real template filled in the way a person would,
a semicolon CSV in Windows-1252, a sheet with no header, and a fake old .xls."""
import datetime as dt
import os
import sys

from openpyxl import Workbook, load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "sheets")
os.makedirs(OUT, exist_ok=True)
TEMPLATE = os.path.join(HERE, "..", "template", "ringgit-log-template.xlsx")
D, T = dt.date, dt.time

# Date, Time, Type, Kind, Description, Amount, Note, Place, Company, In, Out, From, To, km, min, Round, ID
ROWS = [
    (D(2026, 10, 1), T(8, 5), "Travel", "Toll", "PLUS Sungai Buloh", 2.10),
    (D(2026, 10, 1), T(18, 20), "Parking", "Office parking", "Office car park", 9, None, "Menara Ujian", "Metro Parking", T(8, 30), T(18, 15)),
    (D(2026, 10, 2), None, "Travel", "Trip", None, 7.80, "Client visit", None, None, None, None, "Subang Jaya", "Cyberjaya", 33.2, 36, "Yes"),
    (D(2026, 10, 3), T(12, 30), "Food", None, "Nasi kandar", 13.5),
    ("05/10/2026", "1:15 PM", "makan", None, "Kopitiam", "RM 8,50"),
    (D(2026, 10, 6), None, "Parking", "On-site parking", "Mid Valley", 6, None, "Mid Valley Megamall"),
    (D(2026, 10, 7), None, "Gym", None, "Monthly gym", 120),
    (D(2026, 10, 8), None, "Food", None, "Teh tarik", None),
    ("next tuesday", None, "Food", None, "Roti canai", 5),
    (),
    (D(2026, 10, 9), T(7, 40), "Travel", "Fuel", "Petronas Jalan Duta", 50),
    (D(2026, 10, 10), None, "Travel", "Trip", None, 4.6, None, None, None, None, None, "Home", "KLIA"),
    (D(2026, 10, 11), None, None, "Grab", "Grab to office", 15),
]


def filled_template(path):
    wb = load_workbook(TEMPLATE)
    ws = wb["Expenses"]
    for i, row in enumerate(ROWS):
        for j, v in enumerate(row):
            if v is not None:
                ws.cell(row=6 + i, column=1 + j, value=v)
    wb.save(path)


def csv_file(path):
    head = "Date;Time;Type;Kind;Description;Amount (RM);Note;Parking place;Parking company;Time in;Time out;From;To;Distance (km);Travel time (min);Round trip;Record ID"
    lines = [
        head,
        "12/10/2026;08:30;Food;;Café Mahsuri;12,40;;;;;;;;;;;",
        "12/10/2026;;Parking;On-site parking;Pavilion;8;\"Parked 2 h; client lunch\";Pavilion KL;Secure Parking;12:00;14:05;;;;;;",
        "13/10/2026;;Travel;Toll;Batu Tiga toll;2,10;;;;;;;;;;;",
    ]
    with open(path, "wb") as f:
        f.write(("\r\n".join(lines) + "\r\n").encode("cp1252"))


def no_header(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Notes"
    ws.append(["Shopping list"])
    ws.append(["Eggs", 6])
    ws.append(["Milk", 1])
    wb.save(path)


def old_xls(path):
    with open(path, "wb") as f:
        f.write(bytes([0xD0, 0xCF, 0x11, 0xE0, 0xA1, 0xB1, 0x1A, 0xE1]) + b"\0" * 504)


filled_template(os.path.join(OUT, "filled.xlsx"))
csv_file(os.path.join(OUT, "october.csv"))
no_header(os.path.join(OUT, "notes.xlsx"))
old_xls(os.path.join(OUT, "old.xls"))
print("wrote", sorted(os.listdir(OUT)))
