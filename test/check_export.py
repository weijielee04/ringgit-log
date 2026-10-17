"""Checks an exported workbook: schema-valid parts, opens in openpyxl, and prints its rows.
With --edit OUT it also writes an edited copy for the round-trip import test."""
import datetime as dt
import sys
import zipfile

from lxml import etree
from openpyxl import load_workbook

SCHEMAS = "/mnt/skills/public/xlsx/scripts/office/schemas/ISO-IEC29500-4_2016/"
path = sys.argv[1]
z = zipfile.ZipFile(path)
print("parts:", z.namelist())
schema = etree.XMLSchema(etree.parse(SCHEMAS + "sml.xsd"))
bad = 0
for part in ["xl/workbook.xml", "xl/worksheets/sheet1.xml", "xl/styles.xml"]:
    doc = etree.fromstring(z.read(part))
    if schema.validate(doc):
        print(part, "schema-valid")
    else:
        bad += 1
        print(part, "INVALID:", [str(e) for e in schema.error_log][:5])
for part in ["[Content_Types].xml", "_rels/.rels", "xl/_rels/workbook.xml.rels"]:
    etree.fromstring(z.read(part))
print("package xml parses")

wb = load_workbook(path)
ws = wb.active
print("sheet:", ws.title, "dims:", ws.dimensions, "freeze:", ws.freeze_panes, "validations:", len(ws.data_validations.dataValidation))
head = [c.value for c in ws[1]]
print("header:", head)
for row in ws.iter_rows(min_row=2, values_only=True):
    print("  ", [v.isoformat() if isinstance(v, (dt.date, dt.time, dt.datetime)) else v for v in row])
print("number formats row 2:", [c.number_format for c in ws[2]][:8])

if len(sys.argv) > 3 and sys.argv[2] == "--edit":
    col = {h: i for i, h in enumerate(head)}
    for row in ws.iter_rows(min_row=2):
        desc = row[col["Description"]].value
        if desc == "Nasi kandar":
            row[col["Amount (RM)"]].value = 14
        elif desc == "PLUS Sungai Buloh":
            row[col["Description"]].value = "PLUS Sg Buloh (to office)"
        elif desc == "Kopitiam":
            row[col["Date"]].value = dt.datetime(2026, 9, 30)
    n = ws.max_row + 1
    ws.cell(row=n, column=col["Date"] + 1, value=dt.datetime(2026, 10, 12)).number_format = "dd/mm/yyyy"
    ws.cell(row=n, column=col["Type"] + 1, value="Food")
    ws.cell(row=n, column=col["Description"] + 1, value="Mamak")
    ws.cell(row=n, column=col["Amount (RM)"] + 1, value=9.9)
    wb.save(sys.argv[3])
    print("edited copy written:", sys.argv[3])
sys.exit(1 if bad else 0)
