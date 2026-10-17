"""Builds the Ringgit Log monthly import template.

Column headers here are the ones the app's importer and exporter use. Keep them in step with
COLS in src/06-sheets.js.
"""
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = "/home/claude/ringgit-log/template/ringgit-log-template.xlsx"
FONT = "Arial"
INK, INK2, INK3 = "1D181B", "564D52", "72686E"
TINT = {"all": "F1EEF0", "park": "E2F3EC", "trip": "E3EEFB", "id": "F6F4F5"}
BAND = {"all": "E4DFE2", "park": "C9EBDD", "trip": "CFE1F7", "id": "EAE6E8"}
LINE = Side(style="thin", color="D5CED2")
GRID = Border(left=LINE, right=LINE, top=LINE, bottom=LINE)

COLS = [  # header, width, group, number format
    ("Date *", 12, "all", "dd/mm/yyyy"), ("Time", 8, "all", "hh:mm"), ("Type *", 11, "all", "@"),
    ("Kind", 16, "all", "@"), ("Description", 30, "all", "@"), ("Amount (RM) *", 13, "all", "#,##0.00"),
    ("Note", 24, "all", "@"),
    ("Parking place", 22, "park", "@"), ("Parking company", 18, "park", "@"), ("Time in", 9, "park", "hh:mm"), ("Time out", 9, "park", "hh:mm"),
    ("From", 20, "trip", "@"), ("To", 20, "trip", "@"), ("Distance (km)", 13, "trip", "0.0"), ("Travel time (min)", 16, "trip", "0"), ("Round trip", 11, "trip", "@"),
    ("Record ID", 17, "id", "@"),
]
TYPES = "Food,Travel,Parking,Shopping,Bills,Other"
KINDS = "Toll,Fuel,Public transport,E-hailing,Trip,Other travel,Office parking,On-site parking"
FIRST, LAST = 6, 505          # entry rows; row 5 is the example
NCOL = len(COLS)
L = {h.split(" *")[0]: get_column_letter(i + 1) for i, (h, *_rest) in enumerate(COLS)}

def font(**kw):
    base = dict(name=FONT, size=10, color=INK)
    base.update(kw)
    return Font(**base)

wb = Workbook()
ws = wb.active
ws.title = "Expenses"
ws.sheet_properties.tabColor = "D32F3D"

# Row 1: title and running total of the entry rows
ws["A1"] = "Ringgit Log · expenses to import"
ws["A1"].font = font(size=14, bold=True)
ws.row_dimensions[1].height = 24
ws[L["Description"] + "1"] = "Total of your rows (RM)"
ws[L["Description"] + "1"].font = font(bold=True, color=INK2)
ws[L["Description"] + "1"].alignment = Alignment(horizontal="right", vertical="center")
ws[L["Amount (RM)"] + "1"] = f"=SUM({L['Amount (RM)']}{FIRST}:{L['Amount (RM)']}{LAST})"
ws[L["Amount (RM)"] + "1"].font = font(size=12, bold=True)
ws[L["Amount (RM)"] + "1"].number_format = "#,##0.00"

# Row 2: legend naming the cells to edit
ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NCOL)
ws["A2"] = ("Type one expense per row, starting at row 6. Columns marked * are required. Fill the Parking columns only for parking "
            "and the Trip columns only for trips. Row 5 is an example and is skipped when you import. See the How to fill tab.")
ws["A2"].font = font(italic=True, color=INK2)
ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
ws.row_dimensions[2].height = 30

# Row 3: column groups
groups = [("all", "Every expense"), ("park", "Parking only"), ("trip", "Trips only"), ("id", "Leave blank")]
col = 1
for g, label in groups:
    span = sum(1 for c in COLS if c[2] == g)
    if span > 1:
        ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + span - 1)
    cell = ws.cell(row=3, column=col, value=label)
    for c in range(col, col + span):
        ws.cell(row=3, column=c).fill = PatternFill("solid", fgColor=BAND[g])
        ws.cell(row=3, column=c).border = GRID
    cell.font = font(bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    col += span

# Row 4: headers
for i, (head, width, g, fmt) in enumerate(COLS, start=1):
    c = ws.cell(row=4, column=i, value=head)
    c.font = font(bold=True)
    c.fill = PatternFill("solid", fgColor=TINT[g])
    c.border = Border(left=LINE, right=LINE, top=LINE, bottom=Side(style="medium", color=INK3))
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[4].height = 30

# Row 5: one example row in the expected format (the importer skips Record ID = EXAMPLE)
example = {
    "Date": dt.date(2026, 10, 1), "Time": dt.time(18, 40), "Type": "Parking", "Kind": "Office parking",
    "Description": "Office car park", "Amount (RM)": 9.00, "Note": "Example row, skipped on import",
    "Parking place": "Wisma Contoh", "Parking company": "Wilson Parking", "Time in": dt.time(8, 15), "Time out": dt.time(18, 40),
    "Record ID": "EXAMPLE",
}
for i, (head, width, g, fmt) in enumerate(COLS, start=1):
    key = head.split(" *")[0]
    c = ws.cell(row=5, column=i, value=example.get(key))
    c.font = font(italic=True, color=INK3)
    c.fill = PatternFill("solid", fgColor="FAF8F9")
    c.border = GRID
    c.number_format = fmt

# Entry rows: formats and a light grid so the area reads as a form
for r in range(FIRST, LAST + 1):
    for i, (head, width, g, fmt) in enumerate(COLS, start=1):
        c = ws.cell(row=r, column=i)
        c.number_format = fmt
        c.border = GRID
        c.font = font(color=INK3) if g == "id" else font()

def add_dv(dv, letter):
    dv.add(f"{letter}5:{letter}{LAST}")
    ws.add_data_validation(dv)

add_dv(DataValidation(type="list", formula1=f'"{TYPES}"', allow_blank=True, errorStyle="warning", showErrorMessage=True,
                      errorTitle="Type", error="Pick Food, Travel, Parking, Shopping, Bills or Other.",
                      promptTitle="Type", prompt="Food, Travel, Parking, Shopping, Bills or Other", showInputMessage=True), L["Type"])
add_dv(DataValidation(type="list", formula1=f'"{KINDS}"', allow_blank=True, errorStyle="warning", showErrorMessage=True,
                      errorTitle="Kind", error="Pick a kind from the list, or leave it blank.",
                      promptTitle="Kind (optional)", prompt="Travel: Toll, Fuel, Public transport, E-hailing, Trip or Other travel. Parking: Office parking or On-site parking.", showInputMessage=True), L["Kind"])
add_dv(DataValidation(type="list", formula1='"Yes,No"', allow_blank=True, errorStyle="warning", showErrorMessage=True,
                      errorTitle="Round trip", error="Yes or No."), L["Round trip"])
add_dv(DataValidation(type="date", operator="between", formula1="DATE(2000,1,1)", formula2="DATE(2099,12,31)", allow_blank=True,
                      errorStyle="warning", showErrorMessage=True, errorTitle="Date", error="Type the date as day/month/year, for example 13/10/2026."), L["Date"])
for key in ("Time", "Time in", "Time out"):
    add_dv(DataValidation(type="time", operator="between", formula1="0", formula2="0.999988425925926", allow_blank=True,
                          errorStyle="warning", showErrorMessage=True, errorTitle=key, error="Type a time such as 08:30 or 6:40 PM."), L[key])
for key in ("Amount (RM)", "Distance (km)", "Travel time (min)"):
    add_dv(DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                          errorStyle="warning", showErrorMessage=True, errorTitle=key, error="Type a number, for example 12.50."), L[key])

ws.freeze_panes = "B5"
ws.sheet_view.zoomScale = 100

# ---------- How to fill ----------
hw = wb.create_sheet("How to fill")
hw.sheet_properties.tabColor = "2A78D6"
hw.column_dimensions["A"].width = 4
hw.column_dimensions["B"].width = 22
hw.column_dimensions["C"].width = 70
hw.column_dimensions["D"].width = 26
hw.column_dimensions["E"].width = 11
hw["B1"] = "How to fill the template"
hw["B1"].font = font(size=14, bold=True)
steps = [
    "Use the Expenses tab. Type one expense per row, starting at row 6. Row 5 is an example and is skipped when you import.",
    "Date, Type and Amount are required. Everything else is optional.",
    "Type dates as day/month/year, for example 13/10/2026. Times can be 24-hour (18:40) or with AM/PM (6:40 PM).",
    "Parking: set Type to Parking and Kind to Office parking or On-site parking. Add the place, the parking company and the times if you have them.",
    "Toll: set Type to Travel and Kind to Toll. Amount is the toll you paid.",
    "Trips: set Type to Travel and Kind to Trip. Fill From, To, Distance (km) and Travel time (min) for the whole trip. Amount is the toll paid, or 0 if there was none. If you went and came back, set Round trip to Yes and enter the totals for both ways.",
    "In Ringgit Log, open Records and choose Import from Excel, then pick this file. You see every row before anything is saved.",
    "Receipts and screenshots can be added later. Open a record and choose Attach receipt. Parking, toll and trip records without one are marked No receipt yet.",
    "Leave Record ID blank. The app fills it in when you export a month to Excel, so importing an edited export updates those records instead of adding copies.",
]
r = 3
for i, s in enumerate(steps, start=1):
    hw.cell(row=r, column=1, value=i).font = font(bold=True, color=INK3)
    c = hw.cell(row=r, column=2, value=s)
    hw.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    c.font = font()
    c.alignment = Alignment(wrap_text=True, vertical="top")
    hw.row_dimensions[r].height = 28 if len(s) < 120 else 42
    r += 1
r += 1
hw.cell(row=r, column=2, value="Columns").font = font(size=12, bold=True)
r += 1
for i, head in enumerate(["Column", "What to enter", "Example", "Required"], start=2):
    c = hw.cell(row=r, column=i, value=head)
    c.font = font(bold=True)
    c.fill = PatternFill("solid", fgColor=TINT["all"])
    c.border = GRID
r += 1
ref = [
    ("Date", "The day you paid.", "13/10/2026", "Yes"),
    ("Time", "When you paid.", "18:40", ""),
    ("Type", "Food, Travel, Parking, Shopping, Bills or Other.", "Parking", "Yes"),
    ("Kind", "Travel: Toll, Fuel, Public transport, E-hailing, Trip or Other travel. Parking: Office parking or On-site parking.", "Office parking", ""),
    ("Description", "Merchant, place or a short description.", "Nasi lemak, Restoran Ali", ""),
    ("Amount (RM)", "What you paid, in ringgit. For a trip, the toll paid (0 if none).", "9.00", "Yes"),
    ("Note", "Anything you want to remember.", "Claimable", ""),
    ("Parking place", "Building or car park.", "Wisma Contoh", ""),
    ("Parking company", "Operator or council.", "Wilson Parking", ""),
    ("Time in / Time out", "When you parked and left. The app works out the hourly rate.", "08:15 / 18:40", ""),
    ("From / To", "Where a trip started and ended.", "Bangsar South / Cyberjaya", ""),
    ("Distance (km)", "Total distance for the trip.", "66.4", ""),
    ("Travel time (min)", "Total driving time in minutes.", "72", ""),
    ("Round trip", "Yes if you went and came back.", "Yes", ""),
    ("Record ID", "Leave blank. Filled in by the app's Excel export.", "", ""),
]
for row in ref:
    for i, v in enumerate(row, start=2):
        c = hw.cell(row=r, column=i, value=v)
        c.font = font(bold=(i == 2))
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.border = GRID
    hw.row_dimensions[r].height = 28 if len(row[1]) > 60 else 16
    r += 1
hw.sheet_view.showGridLines = False

# ---------- Examples ----------
ex = wb.create_sheet("Examples")
ex.sheet_properties.tabColor = "1BAF7A"
ex["A1"] = "These rows only show the format. Import reads the Expenses tab."
ex["A1"].font = font(italic=True, color=INK2)
for i, (head, width, g, fmt) in enumerate(COLS, start=1):
    c = ex.cell(row=2, column=i, value=head)
    c.font = font(bold=True)
    c.fill = PatternFill("solid", fgColor=TINT[g])
    c.border = GRID
    ex.column_dimensions[get_column_letter(i)].width = width
rows = [
    (dt.date(2026, 10, 1), dt.time(8, 20), "Food", None, "Kopitiam breakfast", 7.50, None, None, None, None, None, None, None, None, None, None),
    (dt.date(2026, 10, 1), dt.time(7, 52), "Travel", "Toll", "LDP Penchala", 2.10, None, None, None, None, None, None, None, None, None, None),
    (dt.date(2026, 10, 1), dt.time(18, 40), "Parking", "Office parking", "Office car park", 9.00, None, "Wisma Contoh", "Wilson Parking", dt.time(8, 15), dt.time(18, 40), None, None, None, None, None),
    (dt.date(2026, 10, 3), dt.time(12, 10), "Parking", "On-site parking", "Street parking", 3.00, None, "Jalan Telawi", "DBKL", dt.time(11, 5), dt.time(12, 10), None, None, None, None, None),
    (dt.date(2026, 10, 4), dt.time(9, 40), "Travel", "Fuel", "Petronas RON95", 70.00, None, None, None, None, None, None, None, None, None, None),
    (dt.date(2026, 10, 5), dt.time(21, 0), "Bills", None, "Unifi broadband", 129.00, "Auto-debit", None, None, None, None, None, None, None, None, None),
    (dt.date(2026, 10, 6), dt.time(9, 10), "Travel", "Trip", "Client visit", 11.00, "Toll both ways", None, None, None, None, "Bangsar South", "Cyberjaya", 66.4, 72, "Yes"),
]
for r_i, row in enumerate(rows, start=3):
    for i, v in enumerate(row, start=1):
        c = ex.cell(row=r_i, column=i, value=v)
        c.font = font()
        c.border = GRID
        c.number_format = COLS[i - 1][3]
ex.freeze_panes = "A3"

wb.active = 0
wb.save(OUT)
print("saved", OUT)
